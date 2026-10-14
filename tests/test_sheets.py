from pathlib import Path

import pytest
from openpyxl import load_workbook

from rsc import books, sheets

TRACKER = Path("/mnt/project-files/rsc-job-system/RSC Job Tracker.xlsx")


@pytest.mark.skipif(not TRACKER.exists(), reason="tracker workbook not present")
def test_import_original_tracker(conn):
    counts = sheets.import_tracker(conn, TRACKER)
    assert counts == {"jobs": 1, "payments": 10, "expenses": 0}, "EXAMPLE rows must be skipped"
    job = books.get_job(conn, 1001)
    assert job["job_id"].startswith("1001-") and job["contract_total"] == 33103 and job["schedule_ok"]
    assert books.next_job_number(conn) == 1002


def test_export_roundtrip(conn, smith, tmp_path):
    books.add_expense(conn, "2026-10-01", "RSC-100", 60, "Fuel")
    out = sheets.export(conn, tmp_path / "out.xlsx")
    wb = load_workbook(out)
    assert wb.sheetnames == ["Jobs", "Payments", "Expenses", "RSC Codes"]
    assert wb["Payments"]["A2"].value == "1001-SMI-1"
    assert wb["Expenses"]["B2"].value == "RSC-100"
