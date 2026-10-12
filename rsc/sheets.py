"""Excel in and out: import the original RSC Job Tracker, export for the accountant."""
import sqlite3
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from . import books

FIRST_ROW = 5  # data starts on row 5 in every tab of the tracker workbook


def _d(value) -> str | None:
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    return str(value) if value else None


def import_tracker(conn: sqlite3.Connection, path: str | Path) -> dict:
    wb = load_workbook(path, data_only=True)
    counts = {"jobs": 0, "payments": 0, "expenses": 0}

    jobs = wb["Jobs"]
    for row in jobs.iter_rows(min_row=FIRST_ROW, values_only=True):
        number, last = row[0], row[1]
        if not last or not isinstance(number, (int, float)):
            continue
        if books.get_job(conn, int(number)):
            continue
        books.create_job(conn, str(last), number=int(number), first_name=row[4] or "", address=row[5] or "",
                         job_type=row[6] or "", status=row[7] or "Estimate", estimate_sent=_d(row[8]),
                         contract_signed=_d(row[9]), contract_total=row[10], notes=row[17] or "")
        counts["jobs"] += 1

    for row in wb["Payments"].iter_rows(min_row=FIRST_ROW, values_only=True):
        if not isinstance(row[0], (int, float)) or row[6] is None:
            continue
        seq = books.add_payment(conn, int(row[0]), float(row[6]), description=row[4] or "",
                                due_when=row[5] or "", notes=row[11] or "")
        pk = conn.execute("SELECT id FROM payments WHERE job_number = ? AND seq = ?", (int(row[0]), seq)).fetchone()[0]
        books.update_payment(conn, pk, invoice_sent=_d(row[7]) or "", date_paid=_d(row[8]) or "",
                             amount_paid=row[9] if row[9] is not None else "")
        counts["payments"] += 1

    for row in wb["Expenses"].iter_rows(min_row=FIRST_ROW, values_only=True):
        if row[1] in (None, "") or row[6] is None or "EXAMPLE" in str(row[9] or ""):
            continue
        books.add_expense(conn, date=_d(row[0]) or "", charge_to=row[1], amount=float(row[6]),
                          vendor=row[3] or "", description=row[4] or "", category=row[5] or "",
                          paid_with=row[7] or "", notes=row[9] or "", source="import")
        counts["expenses"] += 1
    return counts


def export(conn: sqlite3.Connection, path: str | Path) -> Path:
    wb = Workbook()
    head_font = Font(name="Arial", bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="2F5D46")
    money = '$#,##0.00;($#,##0.00);"-"'

    def sheet(title, headers, rows, money_cols=()):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill = head_font, head_fill
        for r in rows:
            ws.append(r)
        for col in money_cols:
            for cell in ws.iter_cols(min_col=col, max_col=col, min_row=2):
                for c in cell:
                    c.number_format = money
        for i, h in enumerate(headers, start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(12, len(h) + 4)
        ws.freeze_panes = "A2"

    jobs = books.list_jobs(conn)
    sheet("Jobs", ["Job #", "Job ID", "Last Name", "First Name", "Address", "Job Type", "Status", "Estimate Sent",
                   "Contract Signed", "Contract Total", "Payments Scheduled", "Paid to Date", "Still Owed",
                   "Job Expenses", "Profit (Contract - Expenses)"],
          [[j["number"], j["job_id"], j["last_name"], j["first_name"], j["address"], j["job_type"], j["status"],
            j["estimate_sent"], j["contract_signed"], j["contract_total"], j["scheduled"], j["paid"], j["owed"],
            j["expenses"], j["profit"]] for j in jobs], money_cols=range(10, 16))
    sheet("Payments", ["Payment ID (QB invoice #)", "Job #", "Client", "What it's for", "Due when", "Amount",
                       "Invoice Sent", "Date Paid", "Amount Paid", "Status"],
          [[p["payment_id"], p["job_number"], p["last_name"], p["description"], p["due_when"], p["amount"],
            p["invoice_sent"], p["date_paid"], p["amount_paid"], p["status"]] for p in books.list_payments(conn)],
          money_cols=(6, 9))
    sheet("Expenses", ["Date", "Charge To", "Goes To", "Vendor", "What was bought", "Category", "Amount",
                       "Paid With", "Source", "Notes"],
          [[e["date"], e["charge_to"], e["goes_to"], e["vendor"], e["description"], e["category"], e["amount"],
            e["paid_with"], e["source"], e["notes"]] for e in books.list_expenses(conn, limit=100000)],
          money_cols=(7,))
    sheet("RSC Codes", ["Code", "What it's for", "QuickBooks account", "Spent to date"],
          [[c["code"], c["description"], c["qb_account"], c["spent"]] for c in books.list_rsc_codes(conn)],
          money_cols=(4,))
    del wb["Sheet"]
    path = Path(path)
    wb.save(path)
    return path
